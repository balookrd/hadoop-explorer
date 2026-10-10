#!/usr/bin/env python3
"""
Скрипт генерации корпоративной презентации PowerPoint (.pptx)
для сервиса Hadoop gRPC Replicator на основе материалов docs/replicator-presentation.md.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Цветовая палитра Hadoop Explorer Platform
BG_COLOR = RGBColor(11, 15, 25)          # #0B0F19 темный глубокий фон
CARD_BG = RGBColor(22, 30, 49)           # #161E31 фон карточек
CARD_BORDER = RGBColor(51, 65, 85)       # #334155 границы
TEXT_WHITE = RGBColor(248, 250, 252)     # #F8FAFC основной текст
TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8 вторичный текст
ACCENT_CYAN = RGBColor(6, 182, 212)      # #06B6D4 акцент циан
ACCENT_INDIGO = RGBColor(99, 102, 241)   # #6366F1 индиго
ACCENT_GREEN = RGBColor(16, 185, 129)    # #10B981 зеленый
ACCENT_AMBER = RGBColor(245, 158, 11)    # #F59E0B янтарный
ACCENT_RED = RGBColor(239, 68, 68)       # #EF4444 красный
TABLE_HEADER_BG = RGBColor(30, 41, 59)   # #1E293B заголовок таблицы
TABLE_ROW_ALT = RGBColor(17, 24, 39)     # #111827 альт. строка

TOTAL_SLIDES = 18

def set_shape_bg(shape, color, border_color=None, border_width=Pt(1)):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()

def create_base_slide(prs, slide_num, category, title, description):
    slide_layout = prs.slide_layouts[6] # пустой макет
    slide = prs.slides.add_slide(slide_layout)

    # 1. Заливка фона во весь слайд
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_shape_bg(bg, BG_COLOR)

    # 2. Верхний колонтитул / заголовок
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Категория
    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_CYAN
    p_cat.space_after = Pt(2)

    # Заголовок слайда
    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_after = Pt(4)

    # Подзаголовок / описание
    if description:
        p_desc = tf.add_paragraph()
        p_desc.text = description
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED

    # 3. Нижний колонтитул
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
    ftf = footer_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    p_foot = ftf.paragraphs[0]
    p_foot.text = "Hadoop Explorer Platform • Hadoop gRPC Replicator"
    p_foot.font.size = Pt(9)
    p_foot.font.color.rgb = RGBColor(100, 116, 139)

    p_num = ftf.add_paragraph()
    p_num.text = f"{slide_num} / {TOTAL_SLIDES}"
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.font.size = Pt(9)
    p_num.font.bold = True
    p_num.font.color.rgb = ACCENT_CYAN

    # Тонкая линия разделителя внизу
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Pt(1))
    set_shape_bg(line, RGBColor(30, 41, 59))

    return slide

def add_card(slide, left, top, width, height, title, items, top_border_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_shape_bg(card, CARD_BG, CARD_BORDER, Pt(1))

    # Верхняя цветная полоска акцента, если задана
    if top_border_color:
        accent_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
        set_shape_bg(accent_strip, top_border_color)

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if title:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.space_after = Pt(8)
        first_item = True
    else:
        first_item = False

    for item in items:
        p = tf.add_paragraph() if (title or not first_item) else tf.paragraphs[0]
        first_item = False
        p.text = f"• {item}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(5)

def add_image_card(slide, left, top, width, height, image_path, caption):
    # Рамка для картинки
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_shape_bg(frame, RGBColor(15, 23, 42), CARD_BORDER, Pt(1))

    if os.path.exists(image_path):
        caption_height = Inches(0.35)
        img_top = top + Inches(0.08)
        img_height = height - caption_height - Inches(0.12)
        img_left = left + Inches(0.08)
        img_width = width - Inches(0.16)

        try:
            slide.shapes.add_picture(image_path, img_left, img_top, width=img_width)
        except Exception as e:
            print(f"Предупреждение: ошибка загрузки {image_path}: {e}")

        # Подпись снизу
        tb = slide.shapes.add_textbox(left, top + height - caption_height, width, caption_height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = caption
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(9)
        p.font.italic = True
        p.font.color.rgb = TEXT_MUTED

def add_table_custom(slide, left, top, width, height, headers, rows, col_widths=None):
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    if col_widths:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = w

    # Заголовок
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = ACCENT_CYAN

    # Строки данных
    for row_idx, row_data in enumerate(rows):
        bg = TABLE_ROW_ALT if row_idx % 2 == 1 else CARD_BG
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_WHITE if col_idx == 0 else TEXT_MUTED

def generate_presentation(output_pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    img_dir = os.path.join(base_dir, "docs", "images", "replicator")
    logo_path = os.path.join(base_dir, "images", "logo_white.png")

    print(f"🚀 Генерация 18 слайдов в {output_pptx_path}...")

    # =========================================================================
    # СЛАЙД 1: ТИТУЛЬНЫЙ
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_shape_bg(bg1, BG_COLOR)

    # Логотип
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(0.8), Inches(0.8), width=Inches(3.2))

    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "HADOOP EXPLORER PLATFORM • ИНФРАСТРУКТУРА ДАННЫХ"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf1.add_paragraph()
    p.text = "Hadoop gRPC Replicator"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf1.add_paragraph()
    p.text = "Высокоскоростная межкластерная репликация HDFS и Hive Metastore с защитой от Split-Brain и иерархическим шейпером WAN"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED

    # 6 карточек преимуществ (2 ряда по 3)
    card_w = Inches(3.7)
    card_h = Inches(1.5)
    y_r1 = Inches(3.6)
    y_r2 = Inches(5.3)

    add_card(slide1, Inches(0.8), y_r1, card_w, card_h, "🚀 Высокая скорость gRPC",
             ["Потоковый стриминг Netty минуя YARN", "Чанки по 4 МБ, сжатие Zstd/LZ4", "Tar-Streaming мелких файлов < 1 МБ"], ACCENT_INDIGO)

    add_card(slide1, Inches(4.8), y_r1, card_w, card_h, "🌐 Иерархический WAN Шейпер",
             ["Hierarchical Token Bucket шейпинг", "Лимиты Global, DC-DC и Cluster", "Рантайм-применение за < 1 секунды"], ACCENT_CYAN)

    add_card(slide1, Inches(8.8), y_r1, card_w, card_h, "🛡️ Disaster Recovery Hub",
             ["Аварийный останов Kill-Switch (0 МБ/с)", "Безопасный откат Unfence без перезаписи", "1-Click Reverse Replication (DC2 ➔ DC1)"], ACCENT_GREEN)

    add_card(slide1, Inches(0.8), y_r2, card_w, card_h, "🏛️ Hive Metastore CDC",
             ["Потоковый захват NOTIFICATION_LOG", "Inotify Lease HA без гонок стримов", "HDFS Federation NameService translation"], ACCENT_AMBER)

    add_card(slide1, Inches(4.8), y_r2, card_w, card_h, "🔐 Enterprise Безопасность",
             ["Kerberos SPNEGO SSO и LDAP", "Keytab изоляция + UGI Proxy User doAs", "Полный аудит в Apache Ranger Logs"], RGBColor(236, 72, 153))

    add_card(slide1, Inches(8.8), y_r2, card_w, card_h, "💻 Реактивный Web UI",
             ["Svelte 5 SPA, мониторинг скорости (⚡)", "Расчет ETA и прогресса в реальном времени", "История запусков и Retention Policy"], RGBColor(139, 92, 246))

    # =========================================================================
    # СЛАЙД 2: ПРОБЛЕМАТИКА DISTCP
    # =========================================================================
    slide2 = create_base_slide(prs, 2, "Предпосылки и Мотивация",
                               "Почему Apache DistCp больше не решает задачи бизнеса?",
                               "Классический стек межкластерного копирования Hadoop (MapReduce DistCp) создает критические риски в Enterprise")

    headers2 = ["Фактор деградации", "Проблема классического DistCp (MapReduce)", "Архитектурный ответ Hadoop Replicator"]
    rows2 = [
        ["Конкуренция за YARN", "DistCp запускает тяжелый MR Job; отбирает ресурсы у бизнес-пайплайнов Spark/Flink", "Zero YARN footprint: независимые легковесные демоны на DataNode"],
        ["Неуправляемый WAN", "Забивает межЦОДную магистраль; нет иерархических квот, деградируют клиентские API", "Hierarchical Token Bucket: рантайм-шейпинг Global, DC-DC и Cluster за < 1 сек"],
        ["Шторм мелких файлов", "Файлы < 1 МБ вызывают дисковый bottleneck и перегрузку NameNode RPC сессиями", "Tar-Streaming на лету: упаковка в виртуальный поток с прямой распаковкой в память"],
        ["Метаданные Hive", "Копирует только сырые файлы HDFS; схемы и партиции требуют ручных выгрузок DDL", "HMS CDC Engine: потоковый захват событий из NOTIFICATION_LOG и авто-накатывание DDL"],
        ["Риск Split-Brain в DR", "Нет сетевого ограждения; случайный запуск после сбоя затирает свежие данные на резерве", "DR Failover Hub: сетевой барьер 0 МБ/с, безопасный Unfence и Reverse Replication"]
    ]
    add_table_custom(slide2, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.7), headers2, rows2,
                     [Inches(2.5), Inches(4.6), Inches(4.633)])

    # =========================================================================
    # СЛАЙД 3: ВЫСОКОУРОВНЕВАЯ АРХИТЕКТУРА
    # =========================================================================
    slide3 = create_base_slide(prs, 3, "Архитектура системы",
                               "Архитектура всей конструкции: Control Plane vs Data Plane",
                               "Строгое разделение управляющего контура и прямой потоковой gRPC-магистрали между агентами")

    add_card(slide3, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.7),
             "⚙️ Control Plane (Orchestrator)",
             [
                 "Технологии: Java 21 LTS, Spring Boot 3.3.4, Spring Data JPA, PostgreSQL / SQLite.",
                 "Порт сервиса: 8005 (REST API, SSE подписки, раздача собранного Svelte 5 SPA).",
                 "Job Scheduler: Встроенный Cron Scheduler (ReplicationScheduler) для регламентных синхронизаций.",
                 "TokenBucketThrottler: Централизованный потокобезопасный координатор сетевых квот.",
                 "Inotify Lease Coordinator: Распределенный эксклюзивный лизинг CDC схем для предотвращения дублирования.",
                 "Split-Brain State Machine: Гарантирует целостность состояний при авариях ЦОД и снимках Snapshot."
             ], ACCENT_INDIGO)

    add_card(slide3, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.7),
             "⚡ Data Plane (Worker & Receiver Agents)",
             [
                 "Технологии: Java 21 LTS, gRPC / Protobuf, Netty, Hadoop Client API.",
                 "Порт gRPC: 50051 (mTLS / TLSv1.3 шифрование канала, Zero-Copy передача).",
                 "Стриминг чанками: Передача файлов блоками по 4 МБ со сквозным контролем хэша SHA-256.",
                 "Прямой HDFS I/O: Чтение из локального HDFS и запись в целевой HDFS через Proxy User doAs.",
                 "HMS Thrift Connector: Локальное обращение к Hive Metastore по LAN (порт 9083 Thrift).",
                 "Автономность: Агенты завершают передачу текущего блока даже при перезапуске Оркестратора."
             ], ACCENT_CYAN)

    # =========================================================================
    # СЛАЙД 4: ТОПОЛОГИЯ И ШЕЙПЕР
    # =========================================================================
    slide4 = create_base_slide(prs, 4, "Управление полосой пропускания",
                               "Топология ЦОД и иерархический шейпер (Hierarchical Token Bucket)",
                               "Многоуровневый контроль полосы WAN исключает деградацию клиентских сервисов компании")

    add_card(slide4, Inches(0.8), Inches(1.9), Inches(4.8), Inches(2.2),
             "🎯 3 Уровня сетевого контроля",
             [
                 "Global WAN Cap: Общий лимит всей инфраструктуры (напр. 120 МБ/с).",
                 "DC-DC WAN Limit: Магистральный канал между ЦОД (напр. DC1 ➔ DC2: 100 МБ/с).",
                 "HDFS-HDFS Limit: Квоты между парами кластеров (напр. 60 МБ/с и 40 МБ/с)."
             ], ACCENT_CYAN)

    add_card(slide4, Inches(0.8), Inches(4.3), Inches(4.8), Inches(2.3),
             "⚡ Формула расчета задержки чанка",
             [
                 "delay = max(delay_global, delay_dc_dc, delay_hdfs_hdfs)",
                 "Воркер засыпает по самому узкому горлышку маршрута.",
                 "Рантайм-применение: изменение лимита оператором в UI вступает в силу за < 1 сек без перезапуска воркеров!"
             ], ACCENT_INDIGO)

    add_image_card(slide4, Inches(5.9), Inches(1.9), Inches(6.6), Inches(4.7),
                   os.path.join(img_dir, "04_topology_bandwidth.png"),
                   "Раздел «Топология ЦОД и Полоса»: шейпер Token Bucket (DC-DC, HDFS-HDFS, Global)")

    # =========================================================================
    # СЛАЙД 5: АУТЕНТИФИКАЦИЯ И RBAC
    # =========================================================================
    slide5 = create_base_slide(prs, 5, "Безопасность и Доступ",
                               "Единый вход (SPNEGO SSO / LDAP) и ролевая модель (RBAC)",
                               "Бесшовная интеграция в корпоративный домен безопасности платформы Hadoop Explorer")

    add_card(slide5, Inches(0.8), Inches(1.9), Inches(4.8), Inches(2.2),
             "🔑 Способы входа в систему",
             [
                 "Kerberos SPNEGO SSO: Бесшовный вход в 1 клик по билету ОС.",
                 "LDAP / Active Directory: Авторизация по корпоративным учеткам.",
                 "Демо-профили: Быстрое переключение тестовых ролей (ADM/RW/RO)."
             ], ACCENT_INDIGO)

    add_card(slide5, Inches(0.8), Inches(4.3), Inches(4.8), Inches(2.3),
             "👥 Ролевая модель (RBAC)",
             [
                 "ADMIN (admin_user): Управление всеми задачами, Kill-Switch, лимитами.",
                 "WRITER (de_user): Создание задач в рамках назначенных квот.",
                 "READER (analyst_user): Режим наблюдателя (Read-Only) без права мутаций."
             ], ACCENT_GREEN)

    add_image_card(slide5, Inches(5.9), Inches(1.9), Inches(6.6), Inches(4.7),
                   os.path.join(img_dir, "01_login_screen.png"),
                   "Экран аутентификации LDAP & Kerberos SSO с профилями быстрого переключения ролей")

    # =========================================================================
    # СЛАЙД 6: HDFS DATA PLANE ОПТИМИЗАЦИИ
    # =========================================================================
    slide6 = create_base_slide(prs, 6, "HDFS Data Plane",
                               "Протокол передачи данных, Zero-Staging и Tar-Streaming",
                               "Инженерные решения для надежной и быстрой передачи петабайтных объемов без мусора в хранилище")

    add_card(slide6, Inches(0.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "📦 Tar-Streaming мелких файлов (< 1 MB)",
             [
                 "Файлы < 1 МБ пакуются в виртуальный TAR-поток на лету в RAM.",
                 "Передаются единым непрерывным gRPC-стримом.",
                 "Приемник на лету распаковывает стрим прямо в HDFS (Zero Disk I/O)."
             ], ACCENT_INDIGO)

    add_card(slide6, Inches(6.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "🔒 Zero-Staging и Атомарная фиксация",
             [
                 "Одиночные файлы пишутся во временный файл targetPath + '._staging_'.",
                 "Потоковый расчет контрольной суммы SHA-256.",
                 "После совпадения хэша — мгновенный атомарный fs.rename()."
             ], ACCENT_CYAN)

    add_card(slide6, Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🗜️ Wire Compression (Zstd / LZ4)",
             [
                 "Почанковое сжатие трафика в канале WAN на лету.",
                 "Экономия 40–80% полосы на CSV, JSON, логах и дампах БД.",
                 "Авто-отключение для Parquet/ORC для экономии CPU узлов."
             ], ACCENT_GREEN)

    add_card(slide6, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🧹 4-уровневый HDFS Garbage Collector",
             [
                 "Реактивная очистка staging-файлов при разрыве соединения.",
                 "Предстартовая очистка перед повторным запуском задачи.",
                 "Периодический демон каждые 15 мин удаляет сироты старше 30 мин."
             ], ACCENT_AMBER)

    # =========================================================================
    # СЛАЙД 7: ГЛАВНЫЙ ДАШБОРД И СОЗДАНИЕ ЗАДАЧ
    # =========================================================================
    slide7 = create_base_slide(prs, 7, "Интерфейс оператора",
                               "Главная панель HDFS Replication и создание задач",
                               "Реактивный мониторинг прогресса, динамический расчет ETA и мастер создания задач")

    add_image_card(slide7, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "02_main_dashboard.png"),
                   "Главная панель: интерактивные фильтры статусов, скорость (⚡ МБ/с), ETA и таблица")

    add_image_card(slide7, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "03_create_job_modal.png"),
                   "Мастер создания задачи: выбор путей, расписания Cron и Kerberos doAs")

    # =========================================================================
    # СЛАЙД 8: ИСТОРИЯ ЗАПУСКОВ И RETENTION
    # =========================================================================
    slide8 = create_base_slide(prs, 8, "Аналитика и Аудит",
                               "История запусков задачи и политика хранения (Retention)",
                               "Полная прозрачность каждого периодического запуска с контролем объема базы данных")

    add_card(slide8, Inches(0.8), Inches(1.9), Inches(4.8), Inches(2.2),
             "📜 Детальный журнал выполнений",
             [
                 "Хронология: номер запуска (#1, #2...), статус и источник (Шедулер/Ручной).",
                 "Тайминги: время старта, финиша и длительность с точностью до секунды.",
                 "Метрики: переданный объем и эффективная скорость передачи (⚡ МБ/с)."
             ], ACCENT_INDIGO)

    add_card(slide8, Inches(0.8), Inches(4.3), Inches(4.8), Inches(2.3),
             "⚙️ Retention Policy в рантайме",
             [
                 "Динамическая настройка глубины хранения (от 1 до 500 запусков).",
                 "Кнопка «Применить лимит» мгновенно удаляет устаревшие записи из БД.",
                 "База данных оркестратора защищена от разрастания при частых запусках."
             ], ACCENT_AMBER)

    add_image_card(slide8, Inches(5.9), Inches(1.9), Inches(6.6), Inches(4.7),
                   os.path.join(img_dir, "05_job_history_modal.png"),
                   "Модальное окно Job Runs History: KPI задачи, журнал выполнений и динамический Retention")

    # =========================================================================
    # СЛАЙД 9: HIVE METASTORE CDC АРХИТЕКТУРА
    # =========================================================================
    slide9 = create_base_slide(prs, 9, "Метаданные Data Lake",
                               "Репликация Hive Metastore: CDC и Inotify Lease HA",
                               "Непрерывная синхронизация баз и таблиц Hive без потери позиции и гонок между воркерами")

    add_card(slide9, Inches(0.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "🏛️ Потоковый CDC (NotificationLog)",
             [
                 "Source Agent автономно опрашивает NOTIFICATION_LOG Hive Metastore.",
                 "События ADD_PARTITION, CREATE_TABLE передаются по gRPC.",
                 "Target Agent применяет DDL локально с гарантией deleteData = false."
             ], ACCENT_INDIGO)

    add_card(slide9, Inches(6.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "🔒 Inotify Lease HA (Без гонок)",
             [
                 "Каждая схема захватывается ровно одним агентом (Lease на 60 сек).",
                 "Авто-продление аренды при рапорте прогресса last_event_id.",
                 "При сбое воркера другой агент пула перехватывает стрим без потерь."
             ], ACCENT_CYAN)

    add_card(slide9, Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🌐 HDFS Federation Mapping",
             [
                 "Трансляция URI в путях партиций: hdfs://ns-dc1/ ➔ hdfs://ns-dc2/.",
                 "Поддержка правил сопоставления федерации (federation-mappings).",
                 "Корректная перелинковка sdLocation на целевой кластер."
             ], ACCENT_GREEN)

    add_card(slide9, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🛡️ Non-ACID Gate & Изоляция",
             [
                 "External и Non-Transactional таблицы реплицируются потоково.",
                 "ACID transactional таблицы безопасно пропускаются (SKIPPED_ACID).",
                 "HDFS саб-джобы переноса скрыты из основного списка репликатора."
             ], ACCENT_AMBER)

    # =========================================================================
    # СЛАЙД 10: HMS REPLICATION ДАШБОРД
    # =========================================================================
    slide10 = create_base_slide(prs, 10, "Интерфейс оператора",
                                "Раздел HMS Replication: мониторинг схем и Re-bootstrap",
                                "Управление непрерывной потоковой CDC-репликацией схем Hive Metastore")

    add_image_card(slide10, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "06_hms_replication_dashboard.png"),
                   "Консоль HMS Replication: статус CDC-стримеров, Event Lag и список схем")

    add_image_card(slide10, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "07_create_hms_modal.png"),
                   "Мастер создания схемы: выбор баз, фильтр таблиц и Reconciliation")

    # =========================================================================
    # СЛАЙД 11: SPLIT-BRAIN ПРОБЛЕМА
    # =========================================================================
    slide11 = create_base_slide(prs, 11, "Disaster Recovery",
                                "Проблема Split-Brain и риск деструктивной перезаписи данных",
                                "Почему простое снятие сетевой блокировки без защитной логики уничтожает данные компании")

    add_card(slide11, Inches(0.8), Inches(1.9), Inches(11.733), Inches(1.3),
             "1. Штатный режим: DC1 (Primary) ➔ DC2 (Standby)",
             ["Данные и метаданные непрерывно синхронизируются из основного дата-центра в резервный."], ACCENT_INDIGO)

    add_card(slide11, Inches(0.8), Inches(3.3), Inches(11.733), Inches(1.3),
             "2. Авария DC1 и переключение трафика на DC2",
             ["DC1 падает. Клиенты и ETL переключаются на DC2. Бизнес пишет свежие данные в DC2. DC1 отстает на всю дельту аварии!"], ACCENT_AMBER)

    add_card(slide11, Inches(0.8), Inches(4.7), Inches(5.7), Inches(1.9),
             "❌ Катастрофа наивного снятия блокировки",
             [
                 "Если старые прямые задачи DC1 ➔ DC2 возобновятся автоматически:",
                 "Старый DC1 затрет или удалит свежие файлы на DC2!",
                 "Итог: Безвозвратная потеря бизнес-данных за время аварии."
             ], ACCENT_RED)

    add_card(slide11, Inches(6.8), Inches(4.7), Inches(5.7), Inches(1.9),
             "✅ Решение Hadoop Replicator",
             [
                 "Снятие изоляции восстанавливает ТОЛЬКО сетевой канал (100 МБ/с).",
                 "Старые прямые задачи остаются STOPPED (защита от запуска).",
                 "Поток разворачивается через Reverse Replication (DC2 ➔ DC1)."
             ], ACCENT_GREEN)

    # =========================================================================
    # СЛАЙД 12: 5-ФАЗНЫЙ РЕГЛАМЕНТ DR
    # =========================================================================
    slide12 = create_base_slide(prs, 12, "Регламент непрерывности бизнеса",
                                "5-фазный регламент Disaster Recovery (RPO → 0, RTO < 5 мин)",
                                "Четкий сквозной алгоритм действий системы и оператора на протяжении всего инцидента")

    headers12 = ["Фаза регламента", "Состояние DC1", "Задачи DC1 ➔ DC2", "Лимит канала", "Действия оператора и системы"]
    rows12 = [
        ["1. Авария DC1 (Kill-Switch)", "DOWN / FENCED", "STOPPED (Cron OFF)", "0 МБ/с", "Кнопка Kill-Switch. Сетевой барьер 0 МБ/с, freeze задач, фиксация Snapshot."],
        ["2. Работа на DR площадке", "OFFLINE", "STOPPED", "0 МБ/с", "Бизнес пишет в DC2. На панели накапливается Delta Lag (непереданные байты/DDL)."],
        ["3. Оживание DC1 (Unfence)", "ALIVE (STANDBY)", "STOPPED (Защищены!)", "100 МБ/с", "Кнопка Снять изоляцию. Восстанавливается ТОЛЬКО сеть. Старые задачи НЕ запускаются!"],
        ["4. Догон дельты (Reverse)", "RECEIVER", "STOPPED", "100 МБ/с", "Кнопка Reverse Replication. Авто-генерация зеркал rev-* (DC2 ➔ DC1) до RPO = 0."],
        ["5. Возврат нагрузки (Failback)", "PRIMARY", "SCHEDULED", "100 МБ/с", "Переключение трафика на DC1. Отзыв зеркал (Отозвать ↩). Возобновление штатного цикла."]
    ]
    add_table_custom(slide12, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.7), headers12, rows12,
                     [Inches(2.4), Inches(1.7), Inches(2.2), Inches(1.2), Inches(4.233)])

    # =========================================================================
    # СЛАЙД 13: DR HUB И ОГРАЖДЕНИЕ
    # =========================================================================
    slide13 = create_base_slide(prs, 13, "Интерфейс Disaster Recovery",
                                "DR Hub, аварийный Kill-Switch и сетевое ограждение",
                                "Наглядный мониторинг топологии и изоляция упавшей площадки в один клик")

    add_image_card(slide13, Inches(0.8), Inches(1.9), Inches(6.2), Inches(4.7),
                   os.path.join(img_dir, "08_disaster_recovery_dashboard.png"),
                   "Консоль DR & Failover Hub: мониторинг доступности ЦОД, потока и суммарного лага дельты")

    add_image_card(slide13, Inches(7.3), Inches(1.9), Inches(5.2), Inches(2.25),
                   os.path.join(img_dir, "09_emergency_kill_switch_modal.png"),
                   "Модальное окно аварийного останова (Kill-Switch)")

    add_image_card(slide13, Inches(7.3), Inches(4.35), Inches(5.2), Inches(2.25),
                   os.path.join(img_dir, "10_disaster_recovery_fenced_state.png"),
                   "Индикация подавленного кластера (ПОДАВЛЕН 🔒 и тревожный баннер)")

    # =========================================================================
    # СЛАЙД 14: СНЯТИЕ ИЗОЛЯЦИИ И REVERSE REPLICATION
    # =========================================================================
    slide14 = create_base_slide(prs, 14, "Интерфейс Disaster Recovery",
                                "Безопасное снятие изоляции (Unfence) и Reverse Replication",
                                "Защита от случайных действий и автоматическая генерация встречных задач синхронизации дельты")

    add_image_card(slide14, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "11_rollback_unfence_modal.png"),
                   "Модальное окно Unfence: снятие сетевого барьера без возобновления старых задач")

    add_image_card(slide14, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.7),
                   os.path.join(img_dir, "12_reverse_replication_modal.png"),
                   "Мастер Reverse Replication: разворот потока данных DC2 ➔ DC1 с подтверждением")

    # =========================================================================
    # СЛАЙД 15: БЕЗОПАСНОСТЬ И RANGER
    # =========================================================================
    slide15 = create_base_slide(prs, 15, "Enterprise Security",
                                "Безопасность, Kerberos Context Isolation и Apache Ranger",
                                "Полное соблюдение корпоративных политик безопасности банковского и телеком-сектора")

    add_card(slide15, Inches(0.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "🔐 Kerberos Proxy User & doAs имперсонация",
             [
                 "Воркер аутентифицируется через системный keytab hdfs-replicator.",
                 "Операции выполняются под UGI автора: createProxyUser(user).doAs(...).",
                 "Пользователь не сможет скопировать данные, к которым нет прямого доступа."
             ], ACCENT_INDIGO)

    add_card(slide15, Inches(6.8), Inches(1.9), Inches(5.7), Inches(2.2),
             "🛡️ Полный аудит в Apache Ranger",
             [
                 "Ranger фиксирует реальное имя создателя задачи репликации.",
                 "Запись в аудите: ugi: ivan_ivanov (auth:PROXY via hdfs-replicator).",
                 "Полная прозрачность для службы информационной безопасности (ИБ)."
             ], ACCENT_GREEN)

    add_card(slide15, Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🔒 Data Transfer Protection (SASL)",
             [
                 "Авто-конфигурация dfs.data.transfer.protection = integrity/privacy.",
                 "Шифрование канала передачи блоков DataNode по сети.",
                 "Исключение сбоев Connection reset на Kerberized-узлах."
             ], ACCENT_AMBER)

    add_card(slide15, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.2),
             "🌐 Сквозной TLSv1.3 / mTLS",
             [
                 "Шифрование управляющих REST API эндпоинтов по HTTPS.",
                 "Шифрование gRPC-магистрали между воркерами и приемниками.",
                 "Поддержка взаимной аутентификации по X.509 сертификатам (mTLS)."
             ], ACCENT_CYAN)

    # =========================================================================
    # СЛАЙД 16: СРАВНЕНИЕ С АЛЬТЕРНАТИВАМИ
    # =========================================================================
    slide16 = create_base_slide(prs, 16, "Конкурентный анализ",
                                "Сравнение с альтернативными решениями на рынке",
                                "Почему Hadoop gRPC Replicator превосходит существующие Open Source и коммерческие аналоги")

    headers16 = ["Возможность / Критерий", "Apache DistCp", "Apache Falcon", "WANdisco Fusion", "Hadoop Replicator"]
    rows16 = [
        ["Нагрузка на YARN", "Высокая (MapReduce)", "Высокая (Oozie/MR)", "Свой демон", "✅ Zero YARN (gRPC)"],
        ["Иерархический шейпер WAN", "Нет", "Нет", "Частичный", "✅ 3 уровня в рантайме (< 1с)"],
        ["Tar-Streaming мелких файлов", "Лимитирован", "Нет", "Задержки", "✅ Потоковый Tar на лету"],
        ["CDC Hive Metastore", "Нет", "Только DDL скрипты", "Сложный плагин", "✅ Inotify Lease HA"],
        ["Защита от Split-Brain в DR", "Ручная (Риск потери)", "Ручная", "Консенсус Paxos", "✅ Fencing + Reverse Hub"],
        ["Веб-интерфейс и мониторинг", "Только YARN UI", "Устаревший UI", "Тяжелый портал", "✅ Svelte 5 SPA"],
        ["Стоимость владения", "Бесплатно", "Архив (EoL)", "$100k+ / год", "✅ Собственная платформа"]
    ]
    add_table_custom(slide16, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.7), headers16, rows16,
                     [Inches(3.3), Inches(2.1), Inches(2.1), Inches(2.1), Inches(2.133)])

    # =========================================================================
    # СЛАЙД 17: ЭКСПЛУАТАЦИЯ И ТЕСТЫ
    # =========================================================================
    slide17 = create_base_slide(prs, 17, "DevOps & Production Readiness",
                                "Эксплуатация, мониторинг Prometheus и Smoke-тестирование",
                                "Полная прозрачность для инфраструктурных команд и дежурной смены")

    add_card(slide17, Inches(0.8), Inches(1.9), Inches(3.7), Inches(4.7),
             "📊 Метрики Prometheus",
             [
                 "Эндпоинт :8005/actuator/prometheus.",
                 "replication_bytes_total — переданный объем.",
                 "active_workers — число активных воркеров.",
                 "replication_transfer_rate_mb_s — скорость.",
                 "hms_replication_event_lag — лаг CDC.",
                 "fenced_clusters_count — число подавленных кластеров под Kill-Switch."
             ], ACCENT_INDIGO)

    add_card(slide17, Inches(4.8), Inches(1.9), Inches(3.7), Inches(4.7),
             "💓 Health Checks & Liveness",
             [
                 "Стандарты облачного развертывания.",
                 "/actuator/health/liveness — статус работы JVM.",
                 "/actuator/health/readiness — готовность к приему трафика.",
                 "Интеграция с Kubernetes Ingress и внешними балансировщиками (HAProxy/Nginx)."
             ], ACCENT_CYAN)

    add_card(slide17, Inches(8.8), Inches(1.9), Inches(3.7), Inches(4.7),
             "🧪 Docker Smoke Tests",
             [
                 "Скрипт ./demo/replicator/run-smoke-tests.sh.",
                 "Развертывание 2 изолированных Hadoop ЦОД.",
                 "Проверка начального Bootstrap данных.",
                 "Проверка потокового CDC партиций Hive.",
                 "Верификация контрольных сумм SHA-256."
             ], ACCENT_GREEN)

    # =========================================================================
    # СЛАЙД 18: ИТОГИ И Q&A
    # =========================================================================
    slide18 = create_base_slide(prs, 18, "Заключение",
                                "Итоги и готовность к демонстрации",
                                "Hadoop gRPC Replicator переводит управление межкластерной репликацией на новый уровень")

    add_card(slide18, Inches(0.8), Inches(1.9), Inches(5.7), Inches(3.2),
             "🎯 Бизнес-эффект",
             [
                 "Нулевое влияние репликации на YARN-очереди компании.",
                 "Защита критического корпоративного WAN-канала от деградации.",
                 "RPO ➔ 0 и RTO < 5 минут при катастрофе основного дата-центра.",
                 "Полная согласованность данных HDFS и схем Hive Metastore."
             ], ACCENT_INDIGO)

    add_card(slide18, Inches(6.8), Inches(1.9), Inches(5.7), Inches(3.2),
             "🛠️ Инженерные преимущества",
             [
                 "Современный стек Java 21 LTS, Spring Boot 3, Netty gRPC.",
                 "Zero-Staging и Tar-Streaming мелких файлов без лишнего I/O.",
                 "Бесшовный откат с защитой от Split-Brain и деструктивной перезаписи.",
                 "Удобный, отзывчивый интерфейс Svelte 5 с темной темой."
             ], ACCENT_CYAN)

    # Финальный блок
    q_card = slide18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.3))
    set_shape_bg(q_card, CARD_BG, ACCENT_GREEN, Pt(1.5))
    q_tb = slide18.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(11.333), Inches(1.0))
    q_tf = q_tb.text_frame
    q_tf.word_wrap = True
    qp1 = q_tf.paragraphs[0]
    qp1.text = "Спасибо за внимание! Вопросы и демонстрация стенда"
    qp1.font.size = Pt(16)
    qp1.font.bold = True
    qp1.font.color.rgb = TEXT_WHITE
    qp1.alignment = PP_ALIGN.CENTER

    qp2 = q_tf.add_paragraph()
    qp2.text = "Готовы перейти к демонстрации живой работы сервиса на локальном демо-стенде (http://localhost:8005)"
    qp2.font.size = Pt(11)
    qp2.font.color.rgb = ACCENT_CYAN
    qp2.alignment = PP_ALIGN.CENTER

    prs.save(output_pptx_path)
    print(f"✅ Презентация успешно сохранена: {output_pptx_path} ({os.path.getsize(output_pptx_path) / 1024:.1f} KB)")

if __name__ == "__main__":
    out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "replicator-presentation.pptx")
    if len(sys.argv) > 1:
        out_path = sys.argv[1]
    generate_presentation(out_path)
